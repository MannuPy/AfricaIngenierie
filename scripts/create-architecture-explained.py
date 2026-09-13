from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


OUT = Path(r"C:\Projet\Projet -ING\architecture-expliquee-africa-ingenierie.docx")
NAVY = "0B2E6D"
BLUE = "1E5BD7"
CORAL = "E9344B"
PALE = "F3F6FB"
PALE_BLUE = "E8F0FF"
INK = "111827"
MUTED = "5B677A"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border(cell, color="D5DEEA", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def cell_text(cell, value, bold=False, color=INK, size=9.1):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(str(value))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    header = t.rows[0]
    for i, value in enumerate(headers):
        cell_text(header.cells[i], value, bold=True, color=WHITE, size=8.8)
        shade(header.cells[i], NAVY)
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_values in rows:
        row = t.add_row()
        no_split(row)
        for i, value in enumerate(row_values):
            cell_text(row.cells[i], value, size=8.7)
            shade(row.cells[i], "FFFFFF" if len(t.rows) % 2 else PALE)
    for row in t.rows:
        no_split(row)
        for cell in row.cells:
            border(cell)
    if widths:
        for row in t.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(10.2)
        r.font.color.rgb = RGBColor.from_string(INK)
        r2 = p.add_run(text[len(bold_prefix):])
        r2.font.name = "Arial"
        r2.font.size = Pt(10.2)
        r2.font.color.rgb = RGBColor.from_string(INK)
    else:
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10.2)
        r.font.color.rgb = RGBColor.from_string(INK)
    return p


def bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.17)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("•  " + text)
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(INK)
    return p


def step(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{number}.  ")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10.2)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    r2 = p.add_run(text)
    r2.font.name = "Arial"
    r2.font.size = Pt(10.2)
    r2.font.color.rgb = RGBColor.from_string(INK)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = "Georgia" if level == 1 else "Arial"
    r.bold = True
    r.font.size = Pt(19 if level == 1 else 12.5)
    r.font.color.rgb = RGBColor.from_string(INK)
    return p


def small_label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(CORAL)
    return p


def footer_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("Africa Ingénierie  |  Explication de l’architecture  |  ")
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def setup(doc):
    s = doc.sections[0]
    s.top_margin = Mm(17)
    s.bottom_margin = Mm(16)
    s.left_margin = Mm(18)
    s.right_margin = Mm(18)
    s.header_distance = Mm(8)
    s.footer_distance = Mm(8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    title_style = doc.styles["Title"]
    title_style.font.color.rgb = RGBColor.from_string(INK)
    title_style_ppr = title_style._element.get_or_add_pPr()
    inherited_border = title_style_ppr.find(qn("w:pBdr"))
    if inherited_border is not None:
        title_style_ppr.remove(inherited_border)
    for name, size in (("Heading 1", 19), ("Heading 2", 13), ("Heading 3", 11)):
        st = doc.styles[name]
        st.font.name = "Georgia" if name == "Heading 1" else "Arial"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(INK)
    footer_number(s.footer.paragraphs[0])


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def main():
    doc = Document()
    setup(doc)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(82)
    r = cover.add_run("AFRICA INGÉNIERIE")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(CORAL)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(16)
    tr = title.add_run("Architecture de la plateforme Africa Ingénierie")
    tr.font.name = "Georgia"
    tr.font.size = Pt(28)
    tr.bold = True
    tr.font.color.rgb = RGBColor.from_string(INK)
    # Word peut ajouter une bordure héritée au style Title ; la couverture
    # doit rester typographique, sans filet directement sous le titre.
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Explication simple pour une personne non technique")
    sr.font.name = "Arial"
    sr.font.size = Pt(13)
    sr.font.color.rgb = RGBColor.from_string(MUTED)

    doc.add_paragraph().paragraph_format.space_after = Pt(26)
    table(doc, ["Ce document explique", "Pourquoi c’est important"], [
        ["Les différentes parties du site et leur rôle", "Savoir où se trouve chaque responsabilité"],
        ["Le parcours d’une information", "Comprendre ce qui se passe lorsqu’une personne consulte ou modifie le site"],
        ["Les choix de sécurité et d’organisation", "Protéger les contenus, les visiteurs et les données"],
    ], widths=[3.15, 3.15])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    rr = p.add_run(f"Version explicative du {date.today().strftime('%d/%m/%Y')}")
    rr.font.name = "Arial"
    rr.font.size = Pt(9)
    rr.font.color.rgb = RGBColor.from_string(MUTED)

    page_break(doc)
    small_label(doc, "À retenir")
    heading(doc, "Le message essentiel")
    body(doc, "La plateforme n’est pas un seul bloc compliqué. Elle ressemble plutôt à une entreprise bien organisée : l’accueil reçoit les visiteurs, les équipes spécialisées travaillent en coulisses, un registre conserve les informations et un coffre protège les fichiers. Chaque partie a une mission claire, ce qui rend le site plus sûr, plus facile à faire évoluer et plus simple à dépanner.")
    body(doc, "Le site public est la vitrine visible par les visiteurs. L’administration est le bureau de gestion où l’équipe rédige, traduit, organise et publie les contenus. Les deux parties communiquent, mais elles ne touchent pas directement aux mêmes éléments : cette séparation protège les données et évite qu’une erreur d’affichage mette en danger la base de données.")
    heading(doc, "Une comparaison avec un immeuble", 2)
    table(doc, ["Dans la plateforme", "Dans l’analogie", "Rôle concret"], [
        ["Site public", "Accueil et vitrine", "Présente les pages, produits, réalisations, formations et contacts"],
        ["Administration", "Bureau des équipes", "Permet de rédiger, modifier, traduire et publier"],
        ["Nginx", "Portier de l’immeuble", "Dirige chaque demande vers le bon service"],
        ["CMS Payload", "Secrétariat central", "Applique les règles et organise les contenus"],
        ["PostgreSQL", "Registre sécurisé", "Conserve les textes, réglages, comptes et historiques"],
        ["MinIO", "Coffre de fichiers", "Conserve les logos, photos et autres médias"],
        ["Journal d’audit", "Cahier de suivi", "Garde la trace des actions importantes"],
    ], widths=[1.7, 1.65, 2.95])

    page_break(doc)
    heading(doc, "Le chemin d’une visite sur le site")
    body(doc, "Lorsqu’un visiteur ouvre une page, il n’a pas besoin de connaître la base de données. Il demande simplement une page. Le système se charge de trouver la bonne information, de vérifier qu’elle est publiée et de renvoyer une présentation lisible.")
    table(doc, ["Étape", "Ce qui se passe", "Pourquoi"], [
        ["1", "Le visiteur ouvre une adresse dans son navigateur.", "Le navigateur est le point d’entrée familier pour l’utilisateur."],
        ["2", "Le portier technique reçoit la demande et reconnaît le site public ou l’administration.", "Les deux espaces ont des responsabilités différentes."],
        ["3", "Le site public demande uniquement les contenus nécessaires au CMS.", "Le site public ne se connecte pas directement à la base."],
        ["4", "Le CMS vérifie ce qui est publié puis consulte le registre ou le coffre de fichiers.", "Une donnée non publiée ne doit pas apparaître par erreur."],
        ["5", "La page est assemblée puis affichée.", "Le visiteur obtient une expérience rapide et cohérente."],
    ], widths=[0.5, 3.5, 2.3])
    body(doc, "En résumé : navigateur → portier → site ou administration → CMS → registre ou coffre de fichiers → réponse affichée.", bold_prefix="En résumé :")
    heading(doc, "Les deux adresses locales", 2)
    table(doc, ["Espace", "Adresse locale", "Public concerné"], [
        ["Site public", "http://127.0.0.1:8080", "Visiteurs et équipe de recette"],
        ["Administration", "http://admin.localhost:8080/admin", "Administrateurs et équipe éditoriale"],
    ], widths=[1.5, 3.0, 1.8])

    page_break(doc)
    heading(doc, "Les grandes parties et leur raison d’être")
    heading(doc, "Le site public", 2)
    body(doc, "C’est la partie visible : accueil, expertises, produits, réalisations, formations, événements, à propos et contact. Elle est conçue pour être consultée sur ordinateur, tablette et téléphone. Elle présente uniquement les contenus que l’équipe a décidé de rendre publics.")
    heading(doc, "L’administration", 2)
    body(doc, "C’est l’espace de travail de l’entreprise. On y gère les textes, les images, les logos, les partenaires, les témoignages, les chiffres clés, les coordonnées, les réseaux sociaux, les utilisateurs et les publications. L’administration sert donc à piloter le site sans modifier le code pour chaque changement de contenu.")
    heading(doc, "Le CMS", 2)
    body(doc, "Le CMS est le responsable de l’organisation. Il sait quels champs sont obligatoires, quel contenu est un brouillon, qui peut publier et quelles versions ont existé. Cette couche évite que chacun écrive directement dans les données sans contrôle.")
    heading(doc, "La base PostgreSQL", 2)
    body(doc, "PostgreSQL est le registre structuré. Il conserve les informations sous forme organisée : un contact, une formation, un produit ou un témoignage sont des fiches avec des champs identifiables. Cela permet de rechercher, filtrer, sauvegarder et restaurer les informations.")
    heading(doc, "Le stockage MinIO", 2)
    body(doc, "MinIO est le coffre réservé aux fichiers lourds : logos, photos, visuels et autres médias. La base conserve la fiche du média et ses informations, tandis que le fichier lui-même reste dans ce stockage spécialisé.")
    heading(doc, "Le service de messagerie", 2)
    body(doc, "Il sert à transmettre les notifications liées aux formulaires, par exemple lorsqu’un visiteur envoie une demande de contact. Les messages restent également consultables dans l’administration afin de ne pas dépendre d’un seul e-mail.")

    page_break(doc)
    heading(doc, "Pourquoi séparer ces responsabilités")
    table(doc, ["Choix", "Raison", "Bénéfice pour l’entreprise"], [
        ["Site public séparé de la base", "Une vitrine ne doit pas avoir un accès direct au registre interne.", "Moins de risques et meilleure maîtrise des accès."],
        ["CMS séparé de l’affichage", "La gestion éditoriale et la présentation n’évoluent pas au même rythme.", "Les équipes peuvent gérer les contenus sans dépendre d’un développeur."],
        ["Base séparée des fichiers", "Les textes structurés et les images n’ont pas les mêmes besoins.", "Sauvegardes, recherches et médias sont plus fiables."],
        ["Réseaux internes Docker", "Les services sensibles ne sont pas exposés sur Internet.", "La surface d’attaque est réduite."],
        ["Versions et journalisation", "Chaque publication est une décision et une action traçable.", "On peut comprendre, vérifier et corriger plus facilement."],
        ["Bilinguisme géré par contenu", "Le français et l’anglais doivent rester liés mais indépendants.", "Chaque page peut être complète dans les deux langues."],
    ], widths=[1.7, 2.75, 1.85])
    heading(doc, "Ce que cela change au quotidien", 2)
    bullet(doc, "Modifier un téléphone ou un lien social ne demande pas de modifier une page de code.")
    bullet(doc, "Un rédacteur peut préparer un contenu sans le rendre public immédiatement.")
    bullet(doc, "Un publicateur peut relire puis publier lorsque la fiche est prête.")
    bullet(doc, "Un administrateur peut gérer les comptes, les réglages et les droits.")
    bullet(doc, "Une image peut être remplacée dans la médiathèque sans toucher au texte de la page.")

    page_break(doc)
    heading(doc, "Comment une modification devient visible")
    body(doc, "La publication fonctionne comme une validation interne. Elle évite qu’une fiche partiellement remplie, mal traduite ou sans image accessible apparaisse immédiatement sur le site public.")
    step(doc, 1, "L’équipe ouvre la bonne rubrique dans l’administration et modifie la fiche.")
    step(doc, 2, "Les champs nécessaires sont complétés : texte, langue, image, lien, statut ou référencement selon le cas.")
    step(doc, 3, "La fiche est enregistrée en brouillon pour conserver le travail sans l’afficher.")
    step(doc, 4, "La fiche est relue. Le publicateur ou l’administrateur vérifie la qualité, les deux langues et les médias.")
    step(doc, 5, "La fiche est publiée. Le site public reçoit l’information actualisée après la revalidation de la page.")
    step(doc, 6, "L’action est inscrite dans l’historique afin de savoir qui a fait quoi et quand.")
    heading(doc, "Les trois rôles", 2)
    table(doc, ["Rôle", "Responsabilité", "Limite volontaire"], [
        ["Éditeur", "Prépare les contenus et les soumet à validation.", "Ne peut pas publier directement."],
        ["Publicateur", "Relit et publie ou archive les contenus métier.", "Ne gère pas les paramètres réservés à l’administration."],
        ["Administrateur", "Gère tous les contenus, utilisateurs, réglages, droits et historiques.", "Les journaux d’audit restent non modifiables."],
    ], widths=[1.25, 3.7, 1.95])

    page_break(doc)
    heading(doc, "Les données importantes du site")
    heading(doc, "Contenus et langues", 2)
    body(doc, "Chaque contenu important existe en français et en anglais lorsque la page doit être bilingue. La traduction n’est pas un simple bouton automatique : chaque titre, paragraphe, libellé, bouton, texte alternatif d’image et information de pied de page doit être renseigné dans la bonne langue.")
    heading(doc, "Images et logos", 2)
    body(doc, "L’administration conserve le fichier dans la médiathèque et la fiche conserve son lien, son texte alternatif et ses informations d’affichage. Cette organisation facilite le remplacement d’un visuel et permet de respecter l’accessibilité.")
    heading(doc, "Témoignages", 2)
    body(doc, "Un témoignage envoyé par le public arrive d’abord en attente. L’administrateur peut le laisser en brouillon, le publier, le retirer de la publication ou l’archiver. Le site public ne présente que les témoignages approuvés, dans la limite prévue par le fonctionnement éditorial.")
    heading(doc, "Chiffres clés", 2)
    body(doc, "Les chiffres clés sont des données administrables. L’équipe peut modifier la valeur, le suffixe, le libellé dans chaque langue et la visibilité. Cela permet de faire évoluer une statistique sans intervention technique.")
    heading(doc, "Contact et coordonnées", 2)
    body(doc, "Les coordonnées, le téléphone, les horaires, l’adresse et les liens de carte sont centralisés afin d’éviter des informations différentes entre l’en-tête, la page Contact et le pied de page.")
    heading(doc, "Réseaux sociaux", 2)
    body(doc, "Les liens WhatsApp, Facebook et LinkedIn sont réglés depuis l’administration. Si un numéro WhatsApp est configuré, le clic ouvre une conversation vers l’entreprise. Les autres liens conduisent directement vers les pages officielles.")

    page_break(doc)
    heading(doc, "La sécurité expliquée simplement")
    body(doc, "La sécurité ne repose pas sur une seule serrure. Elle repose sur plusieurs protections qui se complètent. Si une protection est contournée, les autres limitent encore les conséquences.")
    table(doc, ["Protection", "Explication simple", "Effet recherché"], [
        ["Compte et mot de passe", "Seules les personnes identifiées entrent dans l’administration.", "Éviter les accès anonymes."],
        ["Rôles", "Chaque compte reçoit uniquement les capacités nécessaires.", "Limiter les erreurs et les abus."],
        ["Validation des formulaires", "Les informations reçues sont contrôlées avant d’être enregistrées.", "Éviter des données incohérentes ou dangereuses."],
        ["Nettoyage des textes", "Les contenus saisis sont filtrés avant affichage.", "Empêcher qu’un texte injecte du code dans une page."],
        ["Réseau privé", "La base et les médias restent derrière les services internes.", "Ne pas exposer directement les données sensibles."],
        ["Journal d’audit", "Les actions sensibles sont enregistrées.", "Pouvoir comprendre et vérifier une modification."],
        ["Sauvegardes", "La base et les fichiers sont copiés selon une procédure.", "Pouvoir restaurer après une panne ou une erreur."],
    ], widths=[1.5, 3.2, 2.2])
    heading(doc, "Ce qu’il ne faut jamais faire", 2)
    bullet(doc, "Partager un mot de passe d’administration dans un e-mail ou une conversation non sécurisée.")
    bullet(doc, "Modifier directement la base sans sauvegarde et sans vérifier la requête.")
    bullet(doc, "Exposer PostgreSQL, MinIO ou pgAdmin sur Internet.")
    bullet(doc, "Publier une fiche sans vérifier son contenu dans les deux langues.")

    page_break(doc)
    heading(doc, "La maintenance et les sauvegardes")
    body(doc, "Une plateforme bien gérée n’attend pas une panne pour réfléchir à la récupération. Les données structurées et les fichiers doivent être sauvegardés ensemble, car une page peut dépendre à la fois d’un texte en base et d’un visuel dans le stockage.")
    table(doc, ["Élément", "Ce qui doit être préservé", "Pourquoi"], [
        ["PostgreSQL", "Contenus, comptes, réglages, messages et historiques.", "Sans le registre, les fiches et les droits sont perdus."],
        ["MinIO", "Logos, photos et visuels.", "Une fiche sans son image peut devenir incomplète."],
        ["Secrets", "Clés d’accès et variables d’environnement, selon la procédure interne.", "Ils permettent aux services de communiquer sans les afficher."],
        ["Procédure de restauration", "Un mode d’emploi et un test périodique.", "Une sauvegarde n’est utile que si elle peut être restaurée."],
    ], widths=[1.35, 3.45, 2.1])
    heading(doc, "Le rôle de l’IDE de base de données", 2)
    body(doc, "pgAdmin ou DBeaver permet à une personne autorisée de regarder les tables et d’exécuter des requêtes classiques. C’est un outil d’inspection et d’exploitation technique, pas un remplacement de l’administration éditoriale. Pour une modification de contenu, l’administration reste préférable car elle applique les validations, les droits, la traduction et la journalisation prévues par le projet.")
    body(doc, "Dans l’environnement local, PostgreSQL est accessible seulement sur la machine du développeur. L’IDE web pgAdmin est optionnel et s’ouvre sur le port local 5050. Depuis pgAdmin, l’hôte de la base est `postgres`. Depuis un IDE installé sur le poste, l’hôte est `127.0.0.1` et le port est `5432`.")
    heading(doc, "Quand utiliser quel outil", 2)
    table(doc, ["Besoin", "Outil conseillé", "Raison"], [
        ["Modifier un texte, une image ou un réglage", "Administration Payload", "Les règles métier et la journalisation sont appliquées."],
        ["Consulter une table ou vérifier une donnée", "pgAdmin ou DBeaver", "Vue directe et requêtes SQL."],
        ["Restaurer ou migrer la structure", "Procédure technique contrôlée", "Opération à risque qui doit être préparée."],
    ], widths=[2.25, 1.9, 2.75])

    page_break(doc)
    heading(doc, "Que se passe t il en cas de problème")
    table(doc, ["Symptôme", "Explication probable", "Première vérification"], [
        ["La modification n’apparaît pas sur le site public", "La fiche est encore en brouillon ou la page doit être revalidée.", "Vérifier le statut puis actualiser la page."],
        ["Une image ne s’affiche pas", "Le fichier n’est pas disponible, son lien est incorrect ou son texte alternatif manque.", "Vérifier la médiathèque et la fiche du contenu."],
        ["Une traduction reste en français", "Un champ localisé n’a pas de valeur anglaise.", "Parcourir aussi les boutons, le pied de page, les métadonnées et les textes d’image."],
        ["Une page affiche une erreur", "Un service dépendant peut être arrêté ou une donnée est invalide.", "Consulter l’état des services et l’historique de l’action."],
        ["Un formulaire n’est pas reçu", "La validation, la limite anti-abus ou le service de messagerie peut bloquer l’envoi.", "Vérifier le message dans l’administration et les journaux."],
        ["pgAdmin ne s’ouvre pas", "Le profil `db-tools` n’est pas démarré ou Docker Desktop est arrêté.", "Lancer `pnpm db:ide` puis vérifier le conteneur."],
    ], widths=[2.05, 2.75, 2.1])
    heading(doc, "La bonne méthode de diagnostic", 2)
    step(doc, 1, "Décrire précisément ce qui est visible et l’adresse de la page.")
    step(doc, 2, "Vérifier si le problème concerne une seule langue, une seule fiche ou tout le site.")
    step(doc, 3, "Regarder le statut de publication, la présence du média et les champs obligatoires.")
    step(doc, 4, "Consulter l’historique et les journaux avant de modifier directement les données.")
    step(doc, 5, "Corriger dans l’administration, enregistrer, publier puis refaire le parcours comme un visiteur.")

    page_break(doc)
    heading(doc, "Glossaire pour non techniciens")
    table(doc, ["Mot", "Explication"], [
        ["Administration", "Espace privé où l’équipe gère les contenus et les réglages du site."],
        ["API", "Passage organisé par lequel deux parties du système échangent des informations."],
        ["Base de données", "Registre structuré qui conserve les informations et permet de les retrouver."],
        ["CMS", "Outil de gestion de contenu utilisé pour rédiger, organiser et publier."],
        ["Docker", "Méthode qui fait fonctionner chaque service dans un environnement isolé."],
        ["Média", "Fichier comme une photo, un logo ou un visuel."],
        ["MinIO", "Stockage spécialisé utilisé pour conserver les fichiers médias."],
        ["PostgreSQL", "Système utilisé comme base de données du CMS."],
        ["Publication", "Décision de rendre une fiche visible sur le site public."],
        ["Revalidation", "Actualisation d’une page après une modification publiée."],
        ["Rôle", "Niveau d’autorisation associé à un compte utilisateur."],
        ["Sauvegarde", "Copie de sécurité utilisée pour retrouver les données après un incident."],
    ], widths=[1.55, 5.35])
    heading(doc, "Checklist de compréhension", 2)
    bullet(doc, "Je sais que le site public et l’administration sont deux espaces différents.")
    bullet(doc, "Je sais que les textes sont dans PostgreSQL et les fichiers dans MinIO.")
    bullet(doc, "Je sais qu’une fiche doit être enregistrée et publiée pour devenir visible.")
    bullet(doc, "Je sais que les droits dépendent du rôle du compte.")
    bullet(doc, "Je sais que les actions importantes sont conservées dans un historique.")
    bullet(doc, "Je sais que pgAdmin ou DBeaver est réservé à l’inspection technique autorisée.")
    body(doc, "Conclusion : l’architecture est organisée pour que chaque équipe puisse travailler dans son espace, avec des règles adaptées à son rôle. Cette organisation protège les données, facilite les évolutions et permet de gérer le site au quotidien sans dépendre d’une intervention technique pour chaque contenu.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
