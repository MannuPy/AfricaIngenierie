from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs' / 'deploiement-ovh-ingenierieafrica-final.md'
OUTPUT = ROOT / 'docs' / 'deploiement-ovh-ingenierieafrica-final.docx'


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd')
        properties.append(shading)
    shading.set(qn('w:fill'), fill)


def set_cell_border(cell, color='D9D9D9', size='6'):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        properties.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        properties.append(margins)
    for margin, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = margins.find(qn(f'w:{margin}'))
        if node is None:
            node = OxmlElement(f'w:{margin}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    properties.append(header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Page ')
    run.font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(field)


def clean(text):
    text = re.sub(r'\[([^]]+)\]\([^)]*\)', r'\1', text)
    return text.replace('`', '')


def add_text(paragraph, text, bold=False):
    run = paragraph.add_run(clean(text))
    run.bold = bold
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    return run


def add_table(doc, rows):
    parsed = []
    for line in rows:
        values = [cell.strip() for cell in line.strip().strip('|').split('|')]
        parsed.append(values)
    if len(parsed) < 2:
        return
    headers = parsed[0]
    body = [row for row in parsed[2:] if not all(set(cell) <= {'-', ':', ' '} for cell in row)]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    table.autofit = True
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for index, value in enumerate(headers):
        cell = header_cells[index]
        cell.text = ''
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        add_text(paragraph, value, bold=True).font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(cell, '17365D')
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row_index, row in enumerate(body):
        cells = table.add_row().cells
        for index in range(len(headers)):
            value = row[index] if index < len(row) else ''
            cell = cells[index]
            cell.text = ''
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_text(paragraph, value)
            if row_index % 2 == 1:
                set_cell_shading(cell, 'F3F6FA')
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [('Title', 24), ('Heading 1', 16), ('Heading 2', 12.5), ('Heading 3', 11)]:
        style = doc.styles[name]
        style.font.name = 'Aptos Display'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(5)


def build():
    content = SOURCE.read_text(encoding='utf-8').splitlines()
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(title, 'Déploiement OVH du site ingenierieafrica.com')
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    run = add_text(subtitle, 'Préparation contrôlée, périmètre DNS et procédure de mise en production')
    run.italic = True
    run.font.color.rgb = RGBColor(55, 55, 55)

    values = [
        ('Projet', 'Africa Ingénierie'),
        ('Domaine public', 'https://ingenierieafrica.com/'),
        ('Date', '12 septembre 2026'),
    ]
    for key, value in values:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        add_text(paragraph, f'{key} : ', bold=True)
        add_text(paragraph, value)
    doc.add_paragraph()

    in_code = False
    code_lines = []
    table_lines = []

    def flush_table():
        nonlocal table_lines
        if table_lines:
            add_table(doc, table_lines)
            table_lines = []

    def flush_code():
        nonlocal code_lines
        if code_lines:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.35)
            paragraph.paragraph_format.right_indent = Cm(0.35)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(7)
            properties = paragraph._p.get_or_add_pPr()
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'F2F4F7')
            properties.append(shading)
            run = paragraph.add_run('\n'.join(code_lines))
            run.font.name = 'Consolas'
            run._element.rPr.rFonts.set(qn('w:ascii'), 'Consolas')
            run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')
            run.font.size = Pt(8)
            code_lines = []

    for raw in content:
        line = raw.rstrip()
        if line.startswith('```'):
            flush_table()
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith('|'):
            table_lines.append(line)
            continue
        flush_table()
        if not line.strip():
            continue
        match = re.match(r'^(#{1,3})\s+(.*)$', line)
        if match:
            level = len(match.group(1))
            heading_style = 'Heading 1' if level == 1 else 'Heading 2' if level == 2 else 'Heading 3'
            paragraph = doc.add_paragraph(style=heading_style)
            add_text(paragraph, match.group(2))
            continue
        if line.startswith('- '):
            paragraph = doc.add_paragraph(style='List Bullet')
            add_text(paragraph, line[2:])
            continue
        ordered = re.match(r'^\d+\.\s+(.*)$', line)
        if ordered:
            paragraph = doc.add_paragraph(style='List Number')
            add_text(paragraph, ordered.group(1))
            continue
        paragraph = doc.add_paragraph()
        add_text(paragraph, line)

    flush_table()
    flush_code()

    footer = section.footer.paragraphs[0]
    footer.text = 'Africa Ingénierie | Préparation OVH | ingenierieafrica.com'
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(90, 90, 90)
    add_page_number(section.footer.add_paragraph())

    doc.core_properties.title = 'Déploiement OVH du site ingenierieafrica.com'
    doc.core_properties.subject = 'Périmètre DNS, préparation serveur et procédure de déploiement OVH'
    doc.core_properties.author = 'Africa Ingénierie'
    doc.core_properties.comments = 'Document de préparation de déploiement'
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build()
